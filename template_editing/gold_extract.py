"""只读提取首个Gold Set；行级功能依据人工阅读记录，原话完整保留。"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import time

# A列只作来源标签；下列映射依据B列实际话术，而非标签直接转换。
MAPPINGS = {
    1: (['CONSULT', 'PROBLEM'], '仅腰痛主题标题；缺少完整看诊解释，弱原型'),
    2: (['HOOK', 'CONSULT', 'PROBLEM'], '结果先行的痛点标题；没有完整治疗解释'),
    3: (['HOOK', 'CONSULT', 'PROBLEM'], '患者年龄和无法站直的具体问题，适合作为痛点开场'),
    4: (['PRODUCT_BRIDGE', 'PRODUCT_INTRO', 'PRODUCT_VALUE'], '从这些问题过渡到课程中覆盖的身体部位和内容'),
    5: (['CALL_RESULT', 'PROOF', 'PRODUCT_PROOF'], '购买后一周的学习与应用反馈，包含对话形式的效果陈述'),
    6: (['CALL_RESULT', 'PROOF', 'PRODUCT_PROOF'], '连麦者讲述使用方法后的结果并致谢'),
    7: (['PRODUCT_VALUE', 'PRODUCT_PROOF', 'PRODUCT_INTRO'], '承接学员结果，解释130节课与外部学习成本；没有直接下单动作'),
    8: (['PRODUCT_PROOF', 'PRODUCT_VALUE'], '主播用口碑和外部成本说明价值，不能仅凭反馈标签认定连麦'),
    9: (['PROOF', 'PRODUCT_PROOF', 'PRODUCT_VALUE'], '对比时间/成本和第三天结果，主要是效果证明而非直接CTA'),
    10: (['PRODUCT_PROOF', 'PRODUCT_VALUE'], '零差评与解决问题的口碑话术；保留原文转写疑词'),
    11: (['PRODUCT_INTRO', 'PRODUCT_VALUE', 'CTA_SOFT'], '介绍视频课程、经验来源、学完能力，并软引导带回去'),
    12: (['CALL_RESULT', 'PROOF', 'PRODUCT_BRIDGE'], '仅治疗反馈衔接带货摘要；没有可核实的完整原话，弱原型'),
    13: (['CALL_RESULT', 'PROOF', 'PRODUCT_BRIDGE'], '仅失眠反馈衔接带货摘要；不能从逼单二字推断CTA子类'),
    14: (['PRODUCT_INTRO', 'PRODUCT_VALUE'], '课程节数/问题覆盖和一家人受益，属于商品介绍与价值'),
    15: (['PRODUCT_INTRO', 'PRODUCT_VALUE'], '亲自出镜、动作频次、通俗易懂等教学形式卖点，属于课程'),
    16: (['PRODUCT_INTRO', 'PRODUCT_VALUE'], '后续更新、笔记、内容稀缺和学习成本，未形成限时催单'),
    17: (['CTA_DIRECT', 'PRODUCT_CTA', 'PRODUCT_VALUE', 'CTA_END'], '明确大家去拍，结合持续更新和早买早受益的收束'),
    18: (['PRODUCT_BRIDGE', 'CTA_DIRECT', 'PRODUCT_CTA'], '从想跟老师学转向直接拍课程'),
    19: (['PRODUCT_INTRO', 'PRODUCT_VALUE'], '课程覆盖、不同年龄都能看懂，强调学习门槛'),
    20: (['CTA_SOFT', 'CTA_DIRECT', 'CTA_URGENCY', 'CTA_END', 'PRODUCT_CTA'], '上架后让需要的人去抢，行动召唤并收束')
}


def neutralize(text):
    # 仅做用户确认的两个实体映射，再将实体弱化；不把人名作为检索锚点。
    return text.replace('苗老师', '陶老师').replace('小灵招', '小妙招').replace('陶老师', '老师').replace('小妙招', '这套课程')


def chunks(text, limit=160):
    # 长案例按标点切分，保留全部内容；不截成编码器前256token。
    phrases = re.split(r'(?<=[。！？?!，,；;])', neutralize(text))
    result, current = [], ''
    for phrase in phrases:
        for i in range(0, len(phrase), limit):
            piece = phrase[i:i + limit]
            if len(current) + len(piece) > limit and current:
                result.append(current)
                current = ''
            current += piece
    if current:
        result.append(current)
    return [r for r in result if r.strip()]


def extract(path, output):
    from openpyxl import load_workbook
    started = time.perf_counter()
    path = Path(path).resolve(strict=True)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows, prototypes = [], []
    try:
        for row in workbook['Sheet1'].iter_rows():
            number = row[0].row
            label, text = row[0].value, row[1].value
            if not text:
                continue
            if number not in MAPPINGS:
                raise ValueError('出现未审核的新案例行，须先记录功能映射')
            functions, why = MAPPINGS[number]
            example_id = f'GOLD-Sheet1-B{number}'
            item = dict(example_id=example_id, source_sheet='Sheet1', source_cell=f'B{number}',
                original_type=label, original_text=str(text), mapped_text=str(text).replace('苗老师','陶老师').replace('小灵招','小妙招'),
                functions=functions, mapping_reason=why, evidence_kind='摘要/标题' if number in [1,2,3,12,13] else '话术原文',
                auxiliary_cell=f'D{number}', auxiliary_value=row[3].value if len(row)>3 else None,
                auxiliary_note='保留原值，不将7.14/12.13等擅自解释为秒数或日期')
            rows.append(item)
            for n, query in enumerate(chunks(str(text)), 1):
                prototypes.append(dict(example_id=example_id, prototype_id=f'{example_id}-P{n}', text=query))
        sheets = workbook.sheetnames
    finally:
        workbook.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != before:
        raise RuntimeError('原Excel发生变化')
    result = dict(source_path=str(path), source_sha256=before, source_unchanged=True, sheets=sheets,
        example_count=len(rows), utterance_count=sum(r['evidence_kind']=='话术原文' for r in rows),
        summary_count=sum(r['evidence_kind']=='摘要/标题' for r in rows),
        function_distribution=dict(Counter(f for r in rows for f in r['functions'])),
        extraction_seconds=time.perf_counter()-started, examples=rows, prototypes=prototypes,
        notes=['C列为DISPIMG图片引用，本轮使用A/B列文本，未把图片内容补写成话术',
               '保留源转写疑词小凝胶/小银招/小林子等，不擅自建立生产商品别名',
               '案例里的疗效/费用等为来源话术，非系统核实事实；相似度不是效果保证'])
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result,ensure_ascii=False,indent=2), 'utf-8')
    print(json.dumps({k:result[k] for k in ['example_count','utterance_count','summary_count','function_distribution']},ensure_ascii=False))
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('xlsx', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    extract(args.xlsx, args.output)
